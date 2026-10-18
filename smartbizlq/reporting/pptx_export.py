"""PowerPoint export functionality for SmartBizLQ dashboard reports."""

import logging
import os
import tempfile
from datetime import datetime, timezone
from typing import Any

import pandas as pd

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    Presentation = None
    Inches = None
    Pt = None

logger = logging.getLogger(__name__)


def export_dashboard_to_pptx(
    kpis: dict[str, Any],
    charts: list[tuple[str, bytes]],
    sample_df: pd.DataFrame,
    out_path: str,
) -> str:
    """Export dashboard KPIs, visual charts, and a sample data preview to a PowerPoint presentation.

    Args:
        kpis: Mapping of KPI title to formatted value string (e.g. {"Total Revenue": "$1,200,000"}).
        charts: List of tuples containing (chart_title, png_image_bytes).
        sample_df: DataFrame sample to include in the presentation tabular view.
        out_path: Absolute or relative destination filesystem path for the .pptx file.

    Returns:
        The out_path string of the saved PowerPoint presentation.

    Raises:
        RuntimeError: If python-pptx is not installed in the environment.
        ValueError: If out_path is empty.
    """
    if Presentation is None or Inches is None or Pt is None:
        raise RuntimeError("python-pptx is required for PPTX export. Install python-pptx to enable.")

    if not out_path:
        raise ValueError("out_path must not be empty.")

    out_dir = os.path.dirname(os.path.abspath(out_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    logger.info(
        "Beginning PPTX generation: %d KPIs, %d charts, target='%s'",
        len(kpis),
        len(charts),
        out_path,
    )

    prs = Presentation()

    # 1. Title Slide
    title_layout = prs.slide_layouts[0]
    title_slide = prs.slides.add_slide(title_layout)
    if title_slide.shapes.title:
        title_slide.shapes.title.text = "SmartBizLQ — Analytics Report"
    if len(title_slide.placeholders) > 1:
        gen_time = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S (UTC)")
        title_slide.placeholders[1].text = f"Executive Business Intelligence Report\nGenerated: {gen_time}"

    # 2. Key Performance Indicators Slide
    kpi_layout = prs.slide_layouts[5]  # Title Only
    kpi_slide = prs.slides.add_slide(kpi_layout)
    if kpi_slide.shapes.title:
        kpi_slide.shapes.title.text = "Key Performance Indicators"

    kpi_box = kpi_slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.4), Inches(5.0))
    kpi_tf = kpi_box.text_frame
    kpi_tf.word_wrap = True

    if not kpis:
        p = kpi_tf.paragraphs[0]
        p.text = "No KPI metrics were recorded for this period."
    else:
        for i, (kpi_name, kpi_val) in enumerate(kpis.items()):
            p = kpi_tf.paragraphs[0] if i == 0 else kpi_tf.add_paragraph()
            p.text = f"{kpi_name}: {kpi_val}"
            p.font.size = Pt(18)
            p.font.bold = True
            p.space_after = Pt(12)

    # 3. Chart Slides (using a secure temporary directory for binary assets)
    with tempfile.TemporaryDirectory(prefix="smartbizlq_pptx_") as temp_dir:
        for chart_idx, (title_text, png_bytes) in enumerate(charts):
            chart_slide = prs.slides.add_slide(prs.slide_layouts[5])
            if chart_slide.shapes.title:
                chart_slide.shapes.title.text = title_text or f"Visual {chart_idx + 1}"

            if not png_bytes:
                logger.warning("Empty PNG payload received for chart slide '%s'", title_text)
                continue

            temp_img_path = os.path.join(temp_dir, f"chart_{chart_idx}_{os.getpid()}.png")
            try:
                with open(temp_img_path, "wb") as img_file:
                    img_file.write(png_bytes)

                chart_slide.shapes.add_picture(
                    temp_img_path,
                    Inches(0.6),
                    Inches(1.3),
                    width=Inches(8.8),
                )
                logger.debug("Added chart slide for '%s'", title_text)
            except Exception as chart_err:
                logger.warning(
                    "Failed to insert chart picture for '%s': %s",
                    title_text,
                    chart_err,
                )

    # 4. Sample Data Slide (first 10 rows)
    data_slide = prs.slides.add_slide(prs.slide_layouts[5])
    if data_slide.shapes.title:
        data_slide.shapes.title.text = "Sample Data Extract (First 10 Rows)"

    df_preview = sample_df.head(10) if sample_df is not None else pd.DataFrame()
    data_box = data_slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9.0), Inches(5.5))
    data_tf = data_box.text_frame
    data_tf.word_wrap = True

    if df_preview.empty:
        data_tf.paragraphs[0].text = "No preview rows available."
    else:
        text_table = df_preview.to_string(index=False)
        p = data_tf.paragraphs[0]
        p.text = text_table
        p.font.name = "Consolas"
        p.font.size = Pt(9)

    prs.save(out_path)
    logger.info("Successfully exported PowerPoint report to %s", out_path)
    return out_path
