"""Excel export engine for SmartBizLQ scenario models and KPI summaries with live formulas."""

import logging
import os
from typing import Any

import pandas as pd

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
except ImportError:
    Workbook = None
    Alignment = None
    Border = None
    Font = None
    PatternFill = None
    Side = None
    get_column_letter = None
    DefinedName = None

logger = logging.getLogger(__name__)

# Reusable corporate styles
_NAVY_HEADER = "1F4E79"
_LIGHT_BLUE = "D9E1F2"
_ZEBRA_FILL = "F2F5F9"
_BORDER_COLOR = "D9D9D9"


def _apply_header_style(cell, title: str) -> None:
    """Apply standard navy header formatting to an openpyxl cell."""
    cell.value = title
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color=_NAVY_HEADER, end_color=_NAVY_HEADER, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_thin_borders(cell) -> None:
    """Apply subtle borders to a data cell."""
    thin_side = Side(border_style="thin", color=_BORDER_COLOR)
    cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)


def _autofit_columns(ws, max_len: int = 50) -> None:
    """Adjust worksheet column widths to fit cell contents cleanly."""
    for col in ws.columns:
        first_cell = col[0]
        col_letter = get_column_letter(first_cell.column)
        longest = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > longest:
                longest = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(longest + 4, 12), max_len)


_GBP = '"£"#,##0;[Red]-"£"#,##0'
_GBP2 = '"£"#,##0.00'
_INPUT_FILL = "FFF2CC"  # pale yellow = editable assumption (finance-model convention)

# Inputs sheet row for each model input, in ScenarioInputs field order.
_INPUT_ROWS: dict[str, int] = {
    "price": 2,
    "volume": 3,
    "unit_cost": 4,
    "growth_rate": 5,
    "energy_mwh": 6,
    "wholesale_price_mwh": 7,
    "fixed_costs": 8,
}
_INPUT_FORMATS: dict[str, str] = {
    "price": _GBP2,
    "volume": "#,##0",
    "unit_cost": _GBP2,
    "growth_rate": "0.0%",
    "energy_mwh": "#,##0",
    "wholesale_price_mwh": _GBP2,
    "fixed_costs": _GBP,
}


def export_scenario_workbook(
    base: Any,
    out_path: str,
    cases: dict[str, dict[str, float]] | None = None,
    price_steps: tuple[float, ...] = (0.8, 0.9, 1.0, 1.1, 1.2),
    volume_steps: tuple[float, ...] = (0.8, 0.9, 1.0, 1.1, 1.2),
) -> str:
    """Export the scenario model as an Excel workbook driven entirely by live formulas.

    Sheets:
        Inputs       - yellow cells are editable: base assumptions plus Low/High multipliers.
        Calculations - Low / Base / High columns; every value is a formula on Inputs.
        Outputs      - headline results and variance to Base, all formulas.
        Sensitivity  - operating profit for price x volume; headers and cells are formulas,
                       so the grid re-centres when the base inputs change.

    The formulas mirror ``smartbizlq.analytics.scenarios.compute_scenario`` exactly.

    Args:
        base: A ``ScenarioInputs`` instance (base case).
        out_path: Destination .xlsx path.
        cases: Low/High multipliers; defaults to ``DEFAULT_CASES``.
        price_steps: Price multipliers for sensitivity rows.
        volume_steps: Volume multipliers for sensitivity columns.

    Returns:
        ``out_path``.
    """
    from ..analytics.scenarios import DEFAULT_CASES, INPUT_LABELS

    if Workbook is None:
        raise RuntimeError("openpyxl is required for Excel export. Install openpyxl to enable.")
    if not out_path:
        raise ValueError("out_path must not be empty.")
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)

    cases = DEFAULT_CASES if cases is None else cases
    low_mult = cases.get("Low", {})
    high_mult = cases.get("High", {})

    wb = Workbook()

    # ---------------- Inputs ----------------
    ws_in = wb.active
    ws_in.title = "Inputs"
    for c, h in enumerate(["Assumption", "Base value", "Low multiplier", "High multiplier"], start=1):
        _apply_header_style(ws_in.cell(row=1, column=c), h)
    input_fill = PatternFill(start_color=_INPUT_FILL, end_color=_INPUT_FILL, fill_type="solid")
    for name, row in _INPUT_ROWS.items():
        ws_in.cell(row=row, column=1, value=INPUT_LABELS[name]).font = Font(bold=True)
        cells = [
            ws_in.cell(row=row, column=2, value=float(getattr(base, name))),
            ws_in.cell(row=row, column=3, value=float(low_mult.get(name, 1.0))),
            ws_in.cell(row=row, column=4, value=float(high_mult.get(name, 1.0))),
        ]
        cells[0].number_format = _INPUT_FORMATS[name]
        cells[1].number_format = cells[2].number_format = "0.00x"
        for cell in cells:
            cell.fill = input_fill
        for col in range(1, 5):
            _apply_thin_borders(ws_in.cell(row=row, column=col))
    ws_in.cell(row=10, column=1, value="Yellow cells are inputs. Change them and every other sheet recalculates.").font = Font(italic=True)
    ws_in.cell(row=11, column=1, value="Growth scales sales volume and energy use; price, unit costs and fixed costs are held constant.").font = Font(italic=True)
    for name, row in _INPUT_ROWS.items():
        wb.defined_names.add(DefinedName(name, attr_text=f"Inputs!$B${row}"))
    _autofit_columns(ws_in, max_len=45)

    # ---------------- Calculations ----------------
    ws_c = wb.create_sheet("Calculations")
    case_cols = {"Low": "B", "Base": "C", "High": "D"}
    for c, h in enumerate(["Line item", "Low", "Base", "High"], start=1):
        _apply_header_style(ws_c.cell(row=1, column=c), h)
    # Rows 2-8: the inputs each case actually uses.
    for name, row in _INPUT_ROWS.items():
        ws_c.cell(row=row, column=1, value=INPUT_LABELS[name])
        ws_c[f"B{row}"] = f"=Inputs!B{row}*Inputs!C{row}"
        ws_c[f"C{row}"] = f"=Inputs!B{row}"
        ws_c[f"D{row}"] = f"=Inputs!B{row}*Inputs!D{row}"
        for col in "BCD":
            ws_c[f"{col}{row}"].number_format = _INPUT_FORMATS[name]
    ws_c.cell(row=9, column=1, value="Results").font = Font(bold=True)
    lines = [
        (10, "Revenue", "={c}2*{c}3", _GBP),
        (11, "Variable cost", "={c}4*{c}3", _GBP),
        (12, "Energy cost", "={c}6*{c}7", _GBP),
        (13, "Fixed costs", "={c}8", _GBP),
        (14, "Total cost", "=SUM({c}11:{c}13)", _GBP),
        (15, "Gross margin", "={c}10-{c}11-{c}12", _GBP),
        (16, "Gross margin %", "=IF({c}10<>0,{c}15/{c}10,0)", "0.0%"),
        (17, "Operating profit", "={c}15-{c}13", _GBP),
        (18, "Revenue after growth", "={c}10*(1+{c}5)", _GBP),
        (19, "Total cost after growth", "=({c}11+{c}12)*(1+{c}5)+{c}13", _GBP),
        (20, "Operating profit after growth", "={c}18-{c}19", _GBP),
    ]
    for row, label, tmpl, fmt in lines:
        ws_c.cell(row=row, column=1, value=label).font = Font(bold=row in (15, 17, 20))
        for col in case_cols.values():
            ws_c[f"{col}{row}"] = tmpl.format(c=col)
            ws_c[f"{col}{row}"].number_format = fmt
    for row in range(2, 21):
        for col in range(1, 5):
            _apply_thin_borders(ws_c.cell(row=row, column=col))
    ws_c.column_dimensions["A"].width = 38
    for col in "BCD":
        ws_c.column_dimensions[col].width = 16

    # ---------------- Outputs ----------------
    ws_o = wb.create_sheet("Outputs")
    for c, h in enumerate(["Metric", "Low", "Base", "High", "Low vs Base", "High vs Base"], start=1):
        _apply_header_style(ws_o.cell(row=1, column=c), h)
    out_rows = [(10, "Revenue"), (14, "Total cost"), (15, "Gross margin"), (16, "Gross margin %"),
                (17, "Operating profit"), (20, "Operating profit after growth")]
    for r, (src, label) in enumerate(out_rows, start=2):
        ws_o.cell(row=r, column=1, value=label).font = Font(bold=True)
        fmt = "0.0%" if src == 16 else _GBP
        for col in "BCD":
            ws_o[f"{col}{r}"] = f"=Calculations!{col}{src}"
            ws_o[f"{col}{r}"].number_format = fmt
        ws_o[f"E{r}"] = f"=B{r}-C{r}"
        ws_o[f"F{r}"] = f"=D{r}-C{r}"
        ws_o[f"E{r}"].number_format = ws_o[f"F{r}"].number_format = "+0.0%;-0.0%" if src == 16 else _GBP
        for col in range(1, 7):
            _apply_thin_borders(ws_o.cell(row=r, column=col))
    ws_o.column_dimensions["A"].width = 32
    for col in "BCDEF":
        ws_o.column_dimensions[col].width = 16

    # ---------------- Sensitivity ----------------
    ws_s = wb.create_sheet("Sensitivity")
    ws_s["A1"] = "Operating profit (£): selling price (rows) x sales volume (columns), other inputs at Base"
    ws_s["A1"].font = Font(size=12, bold=True, color=_NAVY_HEADER)
    ws_s["A2"] = "Multiplier"
    ws_s["A3"] = "Price / Volume"
    _apply_header_style(ws_s["B3"], "")
    for j, vm in enumerate(volume_steps):
        col = get_column_letter(3 + j)
        ws_s[f"{col}2"] = vm
        ws_s[f"{col}2"].number_format = "0.00x"
        ws_s[f"{col}3"] = f"=Inputs!$B$3*{col}$2"
        ws_s[f"{col}3"].number_format = "#,##0"
        ws_s[f"{col}3"].font = Font(bold=True)
    for i, pm in enumerate(price_steps):
        r = 4 + i
        ws_s[f"A{r}"] = pm
        ws_s[f"A{r}"].number_format = "0.00x"
        ws_s[f"B{r}"] = f"=Inputs!$B$2*$A{r}"
        ws_s[f"B{r}"].number_format = _GBP2
        ws_s[f"B{r}"].font = Font(bold=True)
        for j in range(len(volume_steps)):
            col = get_column_letter(3 + j)
            cell = ws_s[f"{col}{r}"]
            cell.value = f"=($B{r}-Inputs!$B$4)*{col}$3-Inputs!$B$6*Inputs!$B$7-Inputs!$B$8"
            cell.number_format = _GBP
            _apply_thin_borders(cell)
    ws_s.column_dimensions["A"].width = 16
    for j in range(len(volume_steps) + 1):
        ws_s.column_dimensions[get_column_letter(2 + j)].width = 14

    wb.save(out_path)
    logger.info("Wrote scenario workbook to %s", out_path)
    return out_path


def export_kpi_workbook(kpis: dict[str, Any], df: pd.DataFrame, out_path: str) -> str:
    """Export KPIs, summary metrics, and dataset extract to a formatted Excel workbook.

    Args:
        kpis: Dictionary mapping KPI name to value (e.g. {"Total Revenue": "$1,200,000"}).
        df: Underlying DataFrame for summary data extract.
        out_path: Target filesystem path for the output .xlsx file.

    Returns:
        The out_path string of the saved Excel workbook.

    Raises:
        RuntimeError: If openpyxl is not installed.
        ValueError: If out_path is empty.
    """
    if Workbook is None:
        raise RuntimeError("openpyxl is required for Excel export. Install openpyxl to enable.")

    if not out_path:
        raise ValueError("out_path must not be empty.")

    out_dir = os.path.dirname(os.path.abspath(out_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    logger.info("Generating KPI workbook at: %s", out_path)
    wb = Workbook()

    # Sheet 1: KPI_Summary
    ws_kpi = wb.active
    ws_kpi.title = "KPI_Summary"
    ws_kpi.views.sheetView[0].showGridLines = True

    # Title Banner
    banner = ws_kpi.cell(row=1, column=1, value="SmartBizLQ — Executive KPI Summary")
    banner.font = Font(name="Calibri", size=14, bold=True, color=_NAVY_HEADER)

    headers = ["KPI Metric", "Observed Value", "Category / Description"]
    for c_idx, h in enumerate(headers, start=1):
        _apply_header_style(ws_kpi.cell(row=3, column=c_idx), h)

    current_row = 4
    for kpi_name, kpi_val in kpis.items():
        c_name = ws_kpi.cell(row=current_row, column=1, value=str(kpi_name))
        c_val = ws_kpi.cell(row=current_row, column=2, value=str(kpi_val))
        c_desc = ws_kpi.cell(row=current_row, column=3, value="Dashboard Metric")

        c_name.font = Font(bold=True)
        c_val.alignment = Alignment(horizontal="right")
        for cell in (c_name, c_val, c_desc):
            _apply_thin_borders(cell)
        current_row += 1

    _autofit_columns(ws_kpi)

    # Sheet 2: Data_Profile
    ws_profile = wb.create_sheet(title="Data_Profile")
    ws_profile.views.sheetView[0].showGridLines = True

    p_headers = ["Attribute", "Value"]
    for c_idx, h in enumerate(p_headers, start=1):
        _apply_header_style(ws_profile.cell(row=1, column=c_idx), h)

    row_count = len(df) if df is not None else 0
    col_count = len(df.columns) if df is not None else 0
    missing_count = int(df.isnull().sum().sum()) if df is not None else 0

    profile_items = [
        ("Total Records (Rows)", f"{row_count:,}"),
        ("Total Columns", f"{col_count:,}"),
        ("Total Missing Cells", f"{missing_count:,}"),
        ("Numeric Columns", ", ".join(df.select_dtypes(include=["number"]).columns.tolist()[:8]) if df is not None else "None"),
    ]

    for r_idx, (attr, val) in enumerate(profile_items, start=2):
        c1 = ws_profile.cell(row=r_idx, column=1, value=attr)
        c2 = ws_profile.cell(row=r_idx, column=2, value=val)
        c1.font = Font(bold=True)
        _apply_thin_borders(c1)
        _apply_thin_borders(c2)

    _autofit_columns(ws_profile)

    # Sheet 3: Dataset_Extract
    if df is not None and not df.empty:
        ws_data = wb.create_sheet(title="Dataset_Extract")
        ws_data.views.sheetView[0].showGridLines = True
        ws_data.freeze_panes = "A2"

        # Limit to first 10,000 rows to ensure fast serialization
        df_extract = df.head(10000)
        for c_idx, col_name in enumerate(df_extract.columns, start=1):
            _apply_header_style(ws_data.cell(row=1, column=c_idx), str(col_name))

        for r_idx, row in enumerate(df_extract.itertuples(index=False), start=2):
            for c_idx, cell_value in enumerate(row, start=1):
                cell = ws_data.cell(row=r_idx, column=c_idx, value=cell_value)
                if r_idx % 2 == 0:
                    cell.fill = PatternFill(start_color=_ZEBRA_FILL, end_color=_ZEBRA_FILL, fill_type="solid")
                _apply_thin_borders(cell)

        _autofit_columns(ws_data)

    wb.save(out_path)
    logger.info("Successfully wrote KPI workbook to: %s", out_path)
    return out_path
